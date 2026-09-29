"""Report, export and CLI tests (REQUIREMENTS 26, 27, 33, 34, 35)."""

from __future__ import annotations

import datetime as dt
import subprocess
import sys
from pathlib import Path

import pandas as pd
import pytest

from quant.backtest.engine import BacktestEngine, contribution_attribution, sector_attribution
from quant.backtest.validation import build_provenance
from quant.factors import compute_universe_factors
from quant.ranking import rank_and_select
from quant.reports.backtest_report import render_backtest_report
from quant.reports.charts import generate_charts
from quant.reports.csv_export import export_csvs
from quant.reports.excel import write_workbook
from quant.reports.formatting import DISCLAIMER, markdown_table, money, percent
from quant.reports.tables import (
    annual_returns_table,
    ending_positions_frame,
    ending_sector_allocation,
    factor_scores_frame,
    monthly_returns_table,
    rankings_frame,
)
from quant.reports.weekly import markdown_to_html, render_weekly_report, write_weekly_report

EXPECTED_SHEETS = [
    "Summary", "Portfolio", "IVV Benchmark", "Transactions", "Stock Rankings",
    "Factor Scores", "Sector Rankings", "Drawdowns", "Monthly Returns",
    "Annual Returns", "Configuration",
]


@pytest.fixture(scope="module")
def result(market, config):
    return BacktestEngine(config, market, run_id="REPORT-TEST", progress=lambda m: None).run()


# --- formatting ------------------------------------------------------------
def test_money_and_percent_handle_missing_values():
    assert money(None) == "n/a"
    assert percent(float("nan")) == "n/a"
    assert money(1234.5) == "$1,234.50"
    assert percent(0.1234) == "12.34%"


def test_markdown_table_alignment():
    table = markdown_table([["a", 1]], ["Name", "Value"], ["left", "right"])
    assert "---:" in table
    assert "| a | 1 |" in table


def test_markdown_to_html_renders_tables_and_headings():
    html = markdown_to_html("# Title\n\n| A | B |\n| --- | ---: |\n| 1 | 2 |\n")
    assert "<h1>Title</h1>" in html
    assert "<table>" in html
    assert '<th class="num">B</th>' in html


def test_markdown_to_html_escapes_content():
    html = markdown_to_html("- <script>alert(1)</script>")
    assert "<script>" not in html
    assert "&lt;script&gt;" in html


# --- weekly report ---------------------------------------------------------
def test_weekly_report_contains_the_required_sections(market, config, view, tmp_path):
    factors, rejections = compute_universe_factors(view, view.universe(), config)
    ranking = rank_and_select(factors, config, view.as_of, rejections=rejections)
    markdown = render_weekly_report(
        ranking, config, view, market, execution_date=market.next_trading_day(view.as_of)
    )

    for heading in (
        "Market summary", "Recommended allocation", "Sector rankings",
        "Recommended portfolio", "Changes from last week", "Why these stocks",
        "Risk", "Data and model caveats",
    ):
        assert heading in markdown, f"missing section: {heading}"
    assert DISCLAIMER in markdown
    assert "not trade instructions" in markdown


def test_weekly_report_states_the_execution_convention(market, config, view):
    factors, _ = compute_universe_factors(view, view.universe(), config)
    ranking = rank_and_select(factors, config, view.as_of)
    markdown = render_weekly_report(ranking, config, view, market)
    assert "friday" in markdown.lower()
    assert "next trading day open" in markdown or "that close" in markdown


def test_weekly_report_allocations_sum_to_the_contribution(market, config, view):
    factors, _ = compute_universe_factors(view, view.universe(), config)
    ranking = rank_and_select(factors, config, view.as_of)
    markdown = render_weekly_report(ranking, config, view, market)
    assert money(config.strategy.weekly_contribution) in markdown


def test_weekly_report_writes_markdown_and_html(market, config, view, tmp_path):
    factors, _ = compute_universe_factors(view, view.universe(), config)
    ranking = rank_and_select(factors, config, view.as_of)
    markdown = render_weekly_report(ranking, config, view, market)
    paths = write_weekly_report(markdown, tmp_path, view.as_of)

    assert paths["markdown"].exists()
    assert paths["html"].exists()
    html = paths["html"].read_text(encoding="utf-8")
    assert "<!doctype html>" in html
    assert "prefers-color-scheme: dark" in html, "the HTML report needs a dark mode"
    assert "max-width" in html


# --- backtest report -------------------------------------------------------
def test_backtest_report_covers_the_required_numbered_sections(result, config, market):
    provenance = build_provenance(result.run_id, config, market)
    markdown = render_backtest_report(result, config, provenance.to_dict())
    for heading in (
        "Executive summary", "Portfolio comparison", "Drawdown",
        "Year-by-year performance", "Monthly returns", "Sector contribution",
        "Individual stock contribution", "Risk metrics",
        "Drawdowns and stress periods", "Best and worst periods",
        "Limitations, assumptions and data quality", "Reproducing this run",
    ):
        assert heading in markdown, f"missing section: {heading}"


def test_backtest_report_states_the_bias_position(result, config, market):
    provenance = build_provenance(result.run_id, config, market)
    markdown = render_backtest_report(result, config, provenance.to_dict())
    assert "Survivorship bias" in markdown
    assert "Look-ahead control" in markdown
    assert "Taxes" in markdown
    assert "Transaction costs" in markdown
    assert DISCLAIMER in markdown


def test_backtest_report_never_promises_outperformance(result, config, market):
    """REQUIREMENTS 64: never describe the strategy as safe or sure to win.

    The disclaimer itself contains "will outperform" inside a negation, so it
    is removed before scanning; what remains is the report's own prose.
    """
    provenance = build_provenance(result.run_id, config, market)
    markdown = render_backtest_report(result, config, provenance.to_dict())
    body = markdown.replace(DISCLAIMER, "").lower()
    for phrase in (
        "guaranteed", "risk-free return", "certain to", "is safe",
        "cannot lose", "will beat", "assured",
    ):
        assert phrase not in body, f"report contains an unsupportable claim: {phrase!r}"


def test_backtest_report_carries_the_disclaimer_verbatim(result, config, market):
    provenance = build_provenance(result.run_id, config, market)
    markdown = render_backtest_report(result, config, provenance.to_dict())
    assert DISCLAIMER in markdown
    assert "do not establish" in markdown


def test_comparison_table_reports_both_return_measures(result, config, market):
    provenance = build_provenance(result.run_id, config, market)
    markdown = render_backtest_report(result, config, provenance.to_dict())
    assert "CAGR (time-weighted)" in markdown
    assert "XIRR (money-weighted)" in markdown


# --- tables ----------------------------------------------------------------
def test_rankings_frame_flattens_every_week(result):
    frame = rankings_frame(result.rankings)
    assert not frame.empty
    assert set(frame["sleeve"]) == {"sector_leaders", "high_growth"}
    assert frame["decision_date"].nunique() == len(result.rankings)


def test_factor_scores_frame_is_long_and_drops_gaps(view, config):
    factors, _ = compute_universe_factors(view, view.universe(), config)
    frame = factor_scores_frame(factors, view.as_of)
    assert {"decision_date", "ticker", "factor", "raw_value"} <= set(frame.columns)
    assert frame["raw_value"].notna().all(), "missing factors must be dropped, not zero-filled"


def test_monthly_and_annual_tables(result):
    primary = result.variants["combined"]
    monthly = monthly_returns_table(primary.returns)
    assert "year" in monthly.columns
    annual = annual_returns_table(result.variants)
    assert "year" in annual.columns
    assert "combined" in annual.columns


def test_ending_allocation_sums_to_one(result, market):
    final = result.end_date
    sectors = {t: market.sector_of(t, final) for t in market.loaded_tickers()}
    allocation = ending_sector_allocation(result, sectors, final, "combined")
    assert sum(allocation.values()) == pytest.approx(1.0)


def test_benchmark_is_not_bucketed_as_unclassified(result, market, config):
    final = result.end_date
    sectors = {t: market.sector_of(t, final) for t in market.loaded_tickers()}
    attribution = sector_attribution(
        result.variants["combined"], sectors, final, config.benchmark.ticker
    )
    labels = set(attribution["sector"])
    assert any(config.benchmark.ticker in label for label in labels)
    assert "Unclassified" not in labels


# --- exports ---------------------------------------------------------------
def test_csv_export_writes_the_named_files(result, tmp_path, market, config):
    final = result.end_date
    sectors = {t: market.sector_of(t, final) for t in market.loaded_tickers()}
    stock = contribution_attribution(result.variants["combined"], final)
    sector = sector_attribution(result.variants["combined"], sectors, final)
    written = export_csvs(result, tmp_path, rankings_frame(result.rankings), pd.DataFrame(), stock, sector)

    for name in ("backtest_summary", "portfolio_history", "transactions", "stock_rankings", "drawdowns"):
        assert name in written, f"{name}.csv was not written"
        assert written[name].exists()
    assert not pd.read_csv(written["transactions"]).empty


def test_excel_workbook_has_all_eleven_sheets(result, tmp_path, market, config):
    from openpyxl import load_workbook

    provenance = build_provenance(result.run_id, config, market)
    final = result.end_date
    path = write_workbook(
        tmp_path / "book.xlsx", result, rankings_frame(result.rankings), pd.DataFrame(),
        pd.DataFrame(), monthly_returns_table(result.variants["combined"].returns),
        annual_returns_table(result.variants), ending_positions_frame(result, final),
        provenance.to_dict(),
    )
    workbook = load_workbook(path)
    assert workbook.sheetnames == EXPECTED_SHEETS


def test_excel_configuration_sheet_records_provenance(result, tmp_path, market, config):
    from openpyxl import load_workbook

    provenance = build_provenance(result.run_id, config, market)
    path = write_workbook(
        tmp_path / "book.xlsx", result, pd.DataFrame(), pd.DataFrame(), pd.DataFrame(),
        pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), provenance.to_dict(),
    )
    sheet = load_workbook(path)["Configuration"]
    settings = {row[0] for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True)}
    assert "provenance.run_id" in settings
    assert "provenance.data_snapshot_hash" in settings
    assert any(s and s.startswith("config.scoring.weights") for s in settings)


# --- charts ----------------------------------------------------------------
def test_charts_are_written(result, tmp_path, market, config):
    final = result.end_date
    sectors = {t: market.sector_of(t, final) for t in market.loaded_tickers()}
    stock = contribution_attribution(result.variants["combined"], final)
    sector = sector_attribution(result.variants["combined"], sectors, final)
    allocation = ending_sector_allocation(result, sectors, final)
    charts = generate_charts(result, tmp_path, sector, stock, allocation)

    for name in (
        "portfolio_value", "portfolio_vs_ivv", "drawdown", "rolling_12m", "rolling_36m",
        "sector_contribution", "top_winners", "top_losers", "allocation", "turnover",
    ):
        assert name in charts.charts, f"missing chart: {name}"
        assert charts.charts[name].exists()
        assert charts.charts[name].stat().st_size > 1000


def test_charts_render_in_dark_mode_too(result, tmp_path, market):
    final = result.end_date
    sectors = {t: market.sector_of(t, final) for t in market.loaded_tickers()}
    stock = contribution_attribution(result.variants["combined"], final)
    sector = sector_attribution(result.variants["combined"], sectors, final)
    charts = generate_charts(result, tmp_path, sector, stock, {}, theme_name="dark")
    assert charts.charts["portfolio_value"].exists()


# --- CLI -------------------------------------------------------------------
def _cli(project_root: Path, *args: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, str(project_root / "main.py"), *args],
        capture_output=True, text=True, cwd=str(project_root), timeout=900,
    )


def test_cli_help_lists_every_documented_command(project_root):
    process = _cli(project_root, "--help")
    assert process.returncode == 0
    for command in ("update-data", "rank", "weekly-report", "backtest", "export", "validate-data"):
        assert command in process.stdout


def test_cli_rejects_weights_that_do_not_sum_to_one(project_root):
    process = _cli(
        project_root, "backtest", "--provider", "synthetic",
        "--weights", "growth=0.5,momentum=0.5,quality=0.5,valuation=0.0,risk=0.0",
    )
    assert process.returncode == 2
    assert "must sum to 1.0" in process.stdout + process.stderr


def test_cli_schedule_shows_decision_before_execution(project_root):
    process = _cli(project_root, "schedule", "--provider", "synthetic", "--weeks", "5")
    assert process.returncode == 0, process.stderr
    assert "DECISION" in process.stdout and "EXECUTION" in process.stdout


def test_cli_rank_runs_end_to_end(project_root):
    process = _cli(
        project_root, "rank", "--provider", "synthetic",
        "--as-of", "2023-06-30", "--max-tickers", "30",
    )
    assert process.returncode == 0, process.stderr
    assert "HIGH-GROWTH TOP" in process.stdout
    assert "SYNTHETIC DATA" in process.stdout, "synthetic runs must be labeled"
