"""Excel workbook export (REQUIREMENTS 35).

Eleven sheets, in the order the requirements list them. Formatting is kept
deliberately plain - header styling, frozen panes, number formats and column
widths - because the workbook is a data deliverable that people will sort and
pivot, not a designed document.
"""

from __future__ import annotations

import datetime as dt
import json
from pathlib import Path
from typing import Any, Mapping, Sequence

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..logging_config import get_logger

log = get_logger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="E1E0D9")
HEADER_FONT = Font(bold=True, color="0B0B0B")

CURRENCY_FORMAT = '"$"#,##0.00'
PERCENT_FORMAT = "0.00%"
RATIO_FORMAT = "0.000"

# Column-name fragments that select a number format.
CURRENCY_HINTS = (
    "value", "balance", "contribution", "gain", "amount", "cash", "price",
    "basis", "proceeds", "commission", "slippage", "dividend", "invested",
    "pnl", "cap",
)
PERCENT_HINTS = (
    "return", "drawdown", "volatility", "cagr", "xirr", "yield", "weight",
    "turnover", "outperforming", "margin", "growth", "allocation",
)
RATIO_HINTS = ("sharpe", "sortino", "beta", "score", "ratio", "shares")


def _number_format(column: str) -> str | None:
    lowered = str(column).lower()
    for hint in PERCENT_HINTS:
        if hint in lowered:
            return PERCENT_FORMAT
    for hint in CURRENCY_HINTS:
        if hint in lowered:
            return CURRENCY_FORMAT
    for hint in RATIO_HINTS:
        if hint in lowered:
            return RATIO_FORMAT
    return None


def _write_sheet(writer: pd.ExcelWriter, name: str, frame: pd.DataFrame, index: bool = False) -> None:
    """Write one frame and apply header styling, widths and number formats."""
    sheet_name = name[:31]
    if frame is None or frame.empty:
        frame = pd.DataFrame({"note": [f"No data available for {name}."]})
        index = False
    frame.to_excel(writer, sheet_name=sheet_name, index=index)

    worksheet = writer.sheets[sheet_name]
    offset = 1 if index else 0
    for position, column in enumerate(frame.columns, start=1 + offset):
        cell = worksheet.cell(row=1, column=position)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        values = frame[column]
        header_width = len(str(column))
        try:
            sample = values.astype(str).head(200)
            body_width = int(sample.str.len().max()) if len(sample) else 0
        except (TypeError, ValueError):
            body_width = 12
        width = min(42, max(11, header_width + 2, body_width + 2))
        worksheet.column_dimensions[get_column_letter(position)].width = width

        if pd.api.types.is_numeric_dtype(values):
            number_format = _number_format(column)
            if number_format:
                for row in range(2, len(frame) + 2):
                    worksheet.cell(row=row, column=position).number_format = number_format

    worksheet.freeze_panes = worksheet.cell(row=2, column=1 + offset)


def _config_sheet(config_dict: Mapping[str, Any], provenance: Mapping[str, Any]) -> pd.DataFrame:
    """Flatten the config and provenance into key/value rows."""
    rows: list[dict[str, str]] = []

    def walk(prefix: str, value: Any) -> None:
        if isinstance(value, Mapping):
            for key, inner in value.items():
                walk(f"{prefix}.{key}" if prefix else str(key), inner)
        elif isinstance(value, (list, tuple)):
            rows.append({"setting": prefix, "value": ", ".join(str(v) for v in value)})
        else:
            rows.append({"setting": prefix, "value": str(value)})

    walk("config", config_dict)
    for key in (
        "run_id", "created_at", "software_version", "strategy_version",
        "python_version", "platform", "git_commit", "data_snapshot_hash",
    ):
        if key in provenance:
            rows.append({"setting": f"provenance.{key}", "value": str(provenance[key])})
    universe = provenance.get("universe") or {}
    for key, value in universe.items():
        rows.append({"setting": f"provenance.universe.{key}", "value": str(value)})
    for warning in provenance.get("warnings", [])[:50]:
        rows.append({"setting": "provenance.warning", "value": str(warning)})
    return pd.DataFrame(rows)


def write_workbook(
    path: Path,
    result,
    rankings: pd.DataFrame,
    factor_scores: pd.DataFrame,
    sector_rankings: pd.DataFrame,
    monthly_returns: pd.DataFrame,
    annual_returns: pd.DataFrame,
    positions: pd.DataFrame,
    provenance: Mapping[str, Any],
) -> Path:
    """Write the 11-sheet workbook."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    summary = result.comparison_table().reset_index().rename(columns={"index": "Metric"})

    benchmark_name = "ivv"
    benchmark = result.variants.get(benchmark_name)
    benchmark_frame = (
        benchmark.values.reset_index(drop=True) if benchmark is not None else pd.DataFrame()
    )

    transaction_rows = []
    for name, variant in result.variants.items():
        frame = variant.portfolio.transactions_frame()
        if frame.empty:
            continue
        frame = frame.copy()
        frame.insert(0, "variant", name)
        transaction_rows.append(frame)
    transactions = pd.concat(transaction_rows, ignore_index=True) if transaction_rows else pd.DataFrame()

    drawdown_rows = []
    for name, variant in result.variants.items():
        frame = variant.drawdowns.copy()
        if frame.empty:
            continue
        frame.insert(0, "variant", name)
        drawdown_rows.append(frame)
    drawdowns = pd.concat(drawdown_rows, ignore_index=True) if drawdown_rows else pd.DataFrame()

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _write_sheet(writer, "Summary", summary)
        _write_sheet(writer, "Portfolio", positions)
        _write_sheet(writer, "IVV Benchmark", benchmark_frame)
        _write_sheet(writer, "Transactions", transactions)
        _write_sheet(writer, "Stock Rankings", rankings)
        _write_sheet(writer, "Factor Scores", factor_scores)
        _write_sheet(writer, "Sector Rankings", sector_rankings)
        _write_sheet(writer, "Drawdowns", drawdowns)
        _write_sheet(writer, "Monthly Returns", monthly_returns)
        _write_sheet(writer, "Annual Returns", annual_returns)
        _write_sheet(writer, "Configuration", _config_sheet(result.config.to_dict(), provenance))

    log.info("wrote Excel workbook %s", path)
    return path
